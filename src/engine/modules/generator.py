from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from .theme import Theme
from .content import SlideModel
from typing import List

def hex_to_rgb(hex_str: str) -> RGBColor:
    hex_str = hex_str.lstrip('#')
    if len(hex_str) == 3:
        hex_str = ''.join([c*2 for c in hex_str])
    return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))

def set_slide_background(slide, bg_color_hex: str):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = hex_to_rgb(bg_color_hex)

def set_text_safely(text_frame, text: str, font_name: str, size_pt: int, color_hex: str):
    text_frame.clear()
    p = text_frame.paragraphs[0]
    p.text = text
    p.font.name = font_name
    p.font.size = Pt(size_pt)
    p.font.color.rgb = hex_to_rgb(color_hex)

def set_bullets_safely(text_frame, bullets: List[str], font_name: str, size_pt: int, color_hex: str):
    text_frame.clear()
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.name = font_name
        p.font.size = Pt(size_pt)
        p.font.color.rgb = hex_to_rgb(color_hex)

def generate_pptx(theme: Theme, slides: List[SlideModel], output_path: str):
    """
    Module 3: Generates the .pptx file using python-pptx based on the Theme and SlideModels.
    """
    prs = Presentation()
    
    # Very basic layouts: 0 is usually Title, 1 is Title and Content
    title_layout = prs.slide_layouts[0]
    bullet_layout = prs.slide_layouts[1]
    
    for slide_model in slides:
        if slide_model.kind == 'title':
            slide = prs.slides.add_slide(title_layout)
            set_slide_background(slide, theme.bg_color)
            
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            
            lines = slide_model.content.split('\n')
            title_text = ""
            subtitle_text = ""
            
            if len(lines) > 0:
                title_text = lines[0].strip('# ')
            if len(lines) > 1:
                subtitle_text = lines[1]
                
            set_text_safely(title.text_frame, title_text, theme.title_font, 44, theme.primary_color)
            set_text_safely(subtitle.text_frame, subtitle_text, theme.body_font, 20, theme.text_color)
                
        elif slide_model.kind == 'bullets':
            slide = prs.slides.add_slide(bullet_layout)
            set_slide_background(slide, theme.bg_color)
            
            title = slide.shapes.title
            body = slide.placeholders[1]
            
            lines = slide_model.content.split('\n')
            title_text = ""
            bullets = []
            
            for line in lines:
                if line.startswith('#'):
                    title_text = line.strip('# ')
                elif line.startswith('-'):
                    bullets.append(line.strip('- '))
            
            set_text_safely(title.text_frame, title_text, theme.title_font, 36, theme.primary_color)
            set_bullets_safely(body.text_frame, bullets, theme.body_font, 18, theme.text_color)
            
    prs.save(output_path)
